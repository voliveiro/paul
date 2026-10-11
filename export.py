import io
import re
from datetime import datetime

from docx import Document
from docx.shared import Pt, RGBColor

INLINE_RE = re.compile(r"(\*\*.+?\*\*(?!\*)|\*[^*\s][^*]*?\*)")
HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
BULLET_RE = re.compile(r"^\s*[-*•]\s+(.*)$")
NUMBERED_RE = re.compile(r"^\s*\d+[.)]\s+(.*)$")


def _add_runs(paragraph, text, bold=False):
    """Add text to a paragraph, rendering **bold** and *italic* spans."""
    for part in INLINE_RE.split(text):
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            _add_runs(paragraph, part[2:-2], bold=True)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.bold, run.italic = bold, True
        elif part:
            paragraph.add_run(part).bold = bold


def _add_markdown(doc, text, base_level=2):
    """Render the light markdown the agents produce into docx blocks.

    Heading depth is relative: the shallowest heading in the text maps to base_level.
    """
    depths = [len(m.group(1)) for l in text.splitlines() if (m := HEADING_RE.match(l.strip()))]
    top = min(depths, default=1)
    for line in text.splitlines():
        stripped = line.strip()
        if not stripped or stripped in ("---", "***"):
            continue
        if m := HEADING_RE.match(stripped):
            level = min(base_level + len(m.group(1)) - top, 9)
            doc.add_heading(m.group(2).strip("* "), level)
        elif m := BULLET_RE.match(line):
            _add_runs(doc.add_paragraph(style="List Bullet"), m.group(1))
        elif m := NUMBERED_RE.match(line):
            _add_runs(doc.add_paragraph(style="List Number"), m.group(1))
        else:
            _add_runs(doc.add_paragraph(), stripped)


def build_report(proposal, responses, synthesis, synthesis_mode="standard"):
    """Return a BytesIO containing the simulation as a .docx report."""
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    doc.add_heading("Paul — Policy Simulation Report", 0)
    meta = doc.add_paragraph()
    meta.add_run(datetime.now().strftime("Generated %d %B %Y, %H:%M")).italic = True
    meta.add_run("\nAgencies consulted: " + ", ".join(r["name"] for r in responses))
    mode_label = "Synergies" if synthesis_mode == "synergies" else "Standard (fault lines)"
    meta.add_run(f"\nSynthesis mode: {mode_label}")

    doc.add_heading("Policy Proposal", 1)
    for para in proposal.split("\n"):
        if para.strip():
            doc.add_paragraph(para.strip())

    doc.add_page_break()
    doc.add_heading("Agency Responses", 1)
    for r in responses:
        doc.add_heading(r["name"], 2)
        role = doc.add_paragraph()
        run = role.add_run(r.get("role", ""))
        run.italic = True
        run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
        _add_markdown(doc, r["text"], base_level=3)

    doc.add_page_break()
    doc.add_heading("Orchestrator Synthesis", 1)
    _add_markdown(doc, synthesis, base_level=2)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf
